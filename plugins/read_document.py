import os

from plugins.file_ops import _resolve_safe

# Per call - the whole result goes into the model's context, so a long
# document is read a chunk at a time via `start` (each result says where
# the next chunk begins) rather than all at once.
MAX_CHARS = int(os.environ.get("AGENTIC_READ_MAX_CHARS", "8000"))
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".xml", ".log", ".py", ".js"}


def _pdf_text(path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        pages.append(f"--- page {number} ---\n{(page.extract_text() or '').strip()}")
    return "\n\n".join(pages)


def _docx_text(path) -> str:
    from docx import Document

    doc = Document(str(path))
    lines = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style.startswith("heading") or style == "title":
            lines.append(f"## {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _pptx_text(path) -> str:
    from pptx import Presentation

    slides = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [shape.text_frame.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]
        slides.append(f"--- slide {number} ---\n" + "\n".join(texts))
    return "\n\n".join(slides)


_READERS = {".pdf": _pdf_text, ".docx": _docx_text, ".pptx": _pptx_text}


def read_document(path: str, start: int = 0):
    """Extract the text of a PDF, Word (.docx), PowerPoint (.pptx) or plain
    text file in the sandboxed working directory, MAX_CHARS at a time from
    character offset `start`."""
    try:
        target = _resolve_safe(path)
    except ValueError as exc:
        return {"error": str(exc)}
    if not target.is_file():
        return {"error": f"no such file: {path}"}

    extension = target.suffix.lower()
    try:
        if extension in _READERS:
            text = _READERS[extension](target)
        elif extension in TEXT_EXTENSIONS:
            text = target.read_text(encoding="utf-8", errors="replace")
        else:
            return {"error": f"can't read {extension or 'this kind of'} files - supported: PDF, .docx, .pptx and plain text"}
    except Exception as exc:
        return {"error": f"couldn't read {path}: {type(exc).__name__}: {exc}"}

    if not text.strip():
        return {"path": path, "content": "", "note": "no text found - it may be a scanned document or contain only images"}

    start = max(0, int(start or 0))
    chunk = text[start:start + MAX_CHARS]
    result = {"path": path, "total_chars": len(text), "start": start, "content": chunk}
    if start + MAX_CHARS < len(text):
        result["next_start"] = start + MAX_CHARS
    return result


def register(registry):
    registry.register("read_document", read_document, {
        "name": "read_document",
        "description": (
            "Read a PDF, docx, pptx or text file in the working directory (e.g. one the user sent). "
            "Long files come in chunks: if the result has next_start, call again with start=next_start."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "e.g. 'uploads/report.pdf'"},
                "start": {"type": "integer"},
            },
            "required": ["path"],
        },
    })
