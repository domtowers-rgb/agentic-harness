from docx import Document

from plugins.file_ops import SANDBOX_DIR, _resolve_safe, safe_output_name

MAX_SECTIONS = 100
MAX_ITEMS_PER_SECTION = 200

# Told to the model with every created file: it's delivered for it (as a
# chat attachment, or a download link in the web UI). Without this, a
# model tested live made up its own "download here" link to a local
# sandbox path - useless to someone on their phone.
DELIVERY_NOTE = "This file is sent to the user automatically as an attachment - don't include a link or file path in your reply."


def _sanitize_filename(name: str) -> str:
    return safe_output_name(name, ".docx", "document")


def _as_text_list(value, what: str, index: int):
    if value is None:
        return []
    if isinstance(value, str):
        return [value]
    if not isinstance(value, list):
        raise ValueError(f"section {index}: '{what}' must be a list of strings")
    if len(value) > MAX_ITEMS_PER_SECTION:
        raise ValueError(f"section {index} has too many {what} (max {MAX_ITEMS_PER_SECTION})")
    return [str(item) for item in value]


def create_document(title: str, sections: list, filename: str = None):
    """Create a .docx file in the sandboxed working directory: a title,
    then one optional heading plus paragraphs and/or bullet points per
    entry in `sections`."""
    if not title or not isinstance(title, str):
        return {"error": "title is required"}
    if not isinstance(sections, list):
        return {"error": "sections must be a list of {heading, paragraphs, bullets} objects"}
    if len(sections) > MAX_SECTIONS:
        return {"error": f"too many sections (max {MAX_SECTIONS})"}

    try:
        target = _resolve_safe(_sanitize_filename(filename or title))
    except ValueError as exc:
        return {"error": str(exc)}

    try:
        doc = Document()
        doc.add_heading(title, level=0)
        for i, section in enumerate(sections):
            if not isinstance(section, dict):
                return {"error": f"section {i} must be an object with optional 'heading', 'paragraphs' and 'bullets'"}
            if section.get("heading"):
                doc.add_heading(str(section["heading"]), level=1)
            for paragraph in _as_text_list(section.get("paragraphs"), "paragraphs", i):
                doc.add_paragraph(paragraph)
            for bullet in _as_text_list(section.get("bullets"), "bullets", i):
                doc.add_paragraph(bullet, style="List Bullet")

        target.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(target))
    except ValueError as exc:
        return {"error": str(exc)}
    except Exception as exc:
        return {"error": f"failed to create document: {exc}"}

    path = str(target.relative_to(SANDBOX_DIR))
    # "attachment" marks this as a file to hand back to the user - see
    # _run_tool_calls in agentic_harness/main.py.
    return {"status": "written", "path": path, "attachment": path, "section_count": len(sections), "note": DELIVERY_NOTE}


def register(registry):
    registry.register("create_document", create_document, {
        "name": "create_document",
        "description": (
            "Create a Word (.docx) document - e.g. a handout, worksheet, letter or report - and send it "
            "to the user. It starts with the title; each entry in 'sections' adds an optional heading "
            "followed by its paragraphs and then its bullet points."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Document title, shown at the top"},
                "sections": {
                    "type": "array",
                    "description": "The document's sections, in order",
                    "items": {
                        "type": "object",
                        "properties": {
                            "heading": {"type": "string", "description": "Optional section heading"},
                            "paragraphs": {"type": "array", "items": {"type": "string"}},
                            "bullets": {"type": "array", "items": {"type": "string"}},
                        },
                    },
                },
                "filename": {
                    "type": "string",
                    "description": "Output filename, e.g. 'worksheet.docx'. Defaults to a name derived from the title.",
                },
            },
            "required": ["title", "sections"],
        },
    })
