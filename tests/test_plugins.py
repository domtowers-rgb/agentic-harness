import importlib
import socket
import sys

import pytest


class TestCalculator:
    def test_basic_arithmetic(self):
        from plugins.calculator import calculate
        assert calculate("2 * (3 + 4)") == {"result": 14}

    def test_functions_and_constants(self):
        from plugins.calculator import calculate
        assert calculate("sqrt(16)") == {"result": 4.0}

    def test_rejects_code_injection(self):
        from plugins.calculator import calculate
        result = calculate('__import__("os").system("echo pwned")')
        assert "error" in result

    def test_rejects_attribute_access(self):
        from plugins.calculator import calculate
        result = calculate("(1).__class__")
        assert "error" in result

    def test_rejects_garbage_input(self):
        from plugins.calculator import calculate
        result = calculate("not a valid expression $$$")
        assert "error" in result


class TestCurrentTime:
    def test_defaults_to_utc(self):
        from plugins.current_time import get_current_time
        result = get_current_time()
        assert result["timezone"] == "UTC"
        assert "T" in result["iso"]

    def test_named_timezone(self):
        from plugins.current_time import get_current_time
        result = get_current_time("Europe/London")
        assert result["timezone"] == "Europe/London"
        assert "error" not in result

    def test_unknown_timezone_errors_gracefully(self):
        from plugins.current_time import get_current_time
        result = get_current_time("Not/AZone")
        assert "error" in result


class TestFileOps:
    @pytest.fixture(autouse=True)
    def _sandbox(self, tmp_path, monkeypatch):
        monkeypatch.setenv("AGENTIC_FILES_DIR", str(tmp_path))
        import plugins.file_ops as file_ops
        importlib.reload(file_ops)
        self.file_ops = file_ops
        yield

    def test_write_then_read(self):
        self.file_ops.write_file("note.txt", "hello")
        assert self.file_ops.read_file("note.txt") == {"content": "hello", "truncated": False}

    def test_list_files(self):
        self.file_ops.write_file("a.txt", "x")
        result = self.file_ops.list_files()
        assert "a.txt" in result["entries"]

    def test_write_creates_parent_dirs(self):
        result = self.file_ops.write_file("nested/dir/note.txt", "hi")
        assert result["status"] == "written"
        assert self.file_ops.read_file("nested/dir/note.txt")["content"] == "hi"

    def test_read_missing_file(self):
        result = self.file_ops.read_file("nope.txt")
        assert "error" in result

    def test_list_missing_dir(self):
        result = self.file_ops.list_files("nope")
        assert "error" in result

    @pytest.mark.parametrize("path", ["../../../etc/passwd", "/etc/passwd"])
    def test_path_escape_rejected(self, path):
        # These use "/" as the separator, so they're meaningful traversal
        # attempts on every platform.
        result = self.file_ops.read_file(path)
        assert "error" in result
        assert "sandbox" in result["error"]

    @pytest.mark.skipif(sys.platform != "win32", reason="backslash is only a path separator on Windows")
    def test_windows_backslash_escape_rejected(self):
        # On POSIX this string is just an inert literal filename (backslash
        # isn't a separator there), so it correctly stays inside the sandbox
        # and hits "no such file" instead - not a security gap, just not a
        # meaningful traversal attempt on that platform.
        result = self.file_ops.read_file("..\\..\\secret.txt")
        assert "error" in result
        assert "sandbox" in result["error"]

    def test_write_too_large_rejected(self, monkeypatch):
        monkeypatch.setattr(self.file_ops, "MAX_WRITE_CHARS", 10)
        result = self.file_ops.write_file("big.txt", "x" * 11)
        assert "error" in result


class TestPowerpoint:
    @pytest.fixture(autouse=True)
    def _sandbox(self, tmp_path, monkeypatch):
        monkeypatch.setenv("AGENTIC_FILES_DIR", str(tmp_path))
        import plugins.file_ops as file_ops
        importlib.reload(file_ops)
        import plugins.powerpoint as powerpoint
        importlib.reload(powerpoint)
        self.powerpoint = powerpoint
        self.tmp_path = tmp_path
        yield

    def test_creates_a_real_pptx_with_correct_content(self):
        from pptx import Presentation

        result = self.powerpoint.create_presentation(
            title="Test Deck",
            subtitle="A subtitle",
            slides=[{"title": "Slide One", "bullets": ["first", "second"]}],
        )
        assert result["status"] == "written"
        assert result["slide_count"] == 2

        prs = Presentation(str(self.tmp_path / result["path"]))
        slides = list(prs.slides)
        assert len(slides) == 2
        assert slides[0].shapes.title.text == "Test Deck"
        assert slides[0].placeholders[1].text_frame.text == "A subtitle"
        assert slides[1].shapes.title.text == "Slide One"
        body_text = [p.text for p in slides[1].placeholders[1].text_frame.paragraphs]
        assert body_text == ["first", "second"]

    def test_slide_without_bullets_is_fine(self):
        result = self.powerpoint.create_presentation(title="X", slides=[{"title": "Empty"}])
        assert result["status"] == "written"

    def test_filename_defaults_to_sanitized_title(self):
        result = self.powerpoint.create_presentation(title="My Cool Deck!!", slides=[])
        assert result["path"] == "My_Cool_Deck.pptx"

    def test_missing_title_rejected(self):
        assert "error" in self.powerpoint.create_presentation(title="", slides=[])

    def test_slides_must_be_a_list(self):
        assert "error" in self.powerpoint.create_presentation(title="X", slides="nope")

    def test_slide_entry_must_be_a_dict(self):
        assert "error" in self.powerpoint.create_presentation(title="X", slides=["nope"])

    def test_too_many_slides_rejected(self):
        many = [{"title": f"s{i}"} for i in range(self.powerpoint.MAX_SLIDES + 1)]
        assert "error" in self.powerpoint.create_presentation(title="X", slides=many)

    def test_too_many_bullets_rejected(self):
        bullets = [str(i) for i in range(self.powerpoint.MAX_BULLETS_PER_SLIDE + 1)]
        result = self.powerpoint.create_presentation(title="X", slides=[{"title": "s", "bullets": bullets}])
        assert "error" in result

    def test_filename_cannot_escape_sandbox(self):
        result = self.powerpoint.create_presentation(title="X", slides=[], filename="../../evil")
        # sanitization strips path separators before the sandbox check even
        # runs, so this lands safely inside the sandbox rather than erroring -
        # confirm it did NOT escape.
        assert result["status"] == "written"
        assert (self.tmp_path / result["path"]).resolve().parent == self.tmp_path.resolve()


class TestMemory:
    @pytest.fixture(autouse=True)
    def _sandbox(self, tmp_path, monkeypatch):
        monkeypatch.setenv("AGENTIC_FILES_DIR", str(tmp_path))
        import plugins.file_ops as file_ops
        importlib.reload(file_ops)
        import plugins.memory as memory
        importlib.reload(memory)
        self.memory = memory
        yield

    def test_remember_then_recall(self):
        result = self.memory.remember(key="favorite_color", value="teal")
        assert result == {"status": "saved", "key": "favorite_color"}
        assert self.memory.recall(key="favorite_color") == {"key": "favorite_color", "value": "teal"}

    def test_recall_missing_key_errors(self):
        result = self.memory.recall(key="nope")
        assert "error" in result

    def test_recall_with_no_key_lists_everything(self):
        self.memory.remember(key="a", value="1")
        self.memory.remember(key="b", value="2")
        result = self.memory.recall()
        assert result == {"memories": {"a": "1", "b": "2"}}

    def test_forget_removes_key(self):
        self.memory.remember(key="temp", value="x")
        assert self.memory.forget(key="temp") == {"status": "forgotten", "key": "temp"}
        assert "error" in self.memory.recall(key="temp")

    def test_forget_missing_key_errors(self):
        assert "error" in self.memory.forget(key="nope")

    def test_missing_key_on_remember_rejected(self):
        assert "error" in self.memory.remember(key="", value="x")

    def test_value_too_large_rejected(self):
        result = self.memory.remember(key="k", value="x" * (self.memory.MAX_VALUE_CHARS + 1))
        assert "error" in result

    def test_memory_full_rejects_new_keys_but_allows_updates(self, monkeypatch):
        monkeypatch.setattr(self.memory, "MAX_ENTRIES", 2)
        assert self.memory.remember(key="a", value="1")["status"] == "saved"
        assert self.memory.remember(key="b", value="2")["status"] == "saved"
        # updating an existing key when full is fine - it's not a new entry
        assert self.memory.remember(key="a", value="1-updated")["status"] == "saved"
        # a genuinely new key when full is rejected
        assert "error" in self.memory.remember(key="c", value="3")

    def test_persists_to_a_real_json_file(self, tmp_path):
        self.memory.remember(key="k", value="v")
        assert (tmp_path / "memory.json").is_file()
        # a second "process" (fresh import) reads back what was saved
        import importlib as _importlib
        _importlib.reload(self.memory)
        assert self.memory.recall(key="k") == {"key": "k", "value": "v"}


class TestFetchUrl:
    def test_rejects_bad_scheme(self):
        from plugins.fetch_url import fetch_url
        assert "error" in fetch_url("ftp://example.com")

    def test_rejects_no_hostname(self):
        from plugins.fetch_url import fetch_url
        assert "error" in fetch_url("http://")

    def test_rejects_loopback(self):
        from plugins.fetch_url import fetch_url
        assert "error" in fetch_url("http://127.0.0.1:9/")

    def test_rejects_link_local_metadata_address(self):
        from plugins.fetch_url import fetch_url
        assert "error" in fetch_url("http://169.254.169.254/latest/meta-data/")

    def test_rejects_unresolvable_host(self):
        from plugins.fetch_url import fetch_url
        result = fetch_url("https://this-domain-should-not-exist-xyz123.invalid")
        assert "error" in result

    def test_unresolvable_host_is_not_reported_as_a_private_address(self):
        from plugins.fetch_url import fetch_url
        result = fetch_url("https://this-domain-should-not-exist-xyz123.invalid")
        assert "could not look up" in result["error"]
        assert "private" not in result["error"]

    def test_temporary_dns_failure_is_retried_once(self, monkeypatch):
        import plugins.fetch_url as fetch_module
        monkeypatch.setattr(fetch_module, "DNS_RETRY_DELAY_SECONDS", 0)
        calls = []

        def flaky(host, *args, **kwargs):
            calls.append(host)
            if len(calls) == 1:
                raise socket.gaierror(socket.EAI_AGAIN, "Temporary failure in name resolution")
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.215.14", 0))]

        monkeypatch.setattr(socket, "getaddrinfo", flaky)
        assert fetch_module._resolve_pinned_ip("example.org") == ("93.184.215.14", None)
        assert len(calls) == 2

    def test_unknown_host_is_not_retried(self, monkeypatch):
        import plugins.fetch_url as fetch_module
        calls = []

        def missing(host, *args, **kwargs):
            calls.append(host)
            raise socket.gaierror(socket.EAI_NONAME, "Name or service not known")

        monkeypatch.setattr(socket, "getaddrinfo", missing)
        ip, error = fetch_module._resolve_pinned_ip("nope.invalid")
        assert ip is None and "could not look up" in error
        assert len(calls) == 1

    def test_fetches_real_public_url(self):
        from plugins.fetch_url import fetch_url
        result = fetch_url("https://example.com")
        if "error" in result:
            pytest.skip(f"network unavailable: {result['error']}")
        assert result["status"] == 200
        assert len(result["content"]) > 0

    def test_dns_resolved_only_once(self, monkeypatch):
        calls = []
        real_getaddrinfo = socket.getaddrinfo

        def spy(host, *args, **kwargs):
            calls.append(host)
            return real_getaddrinfo(host, *args, **kwargs)

        monkeypatch.setattr(socket, "getaddrinfo", spy)
        from plugins.fetch_url import fetch_url
        result = fetch_url("https://example.com")
        if "error" in result:
            pytest.skip(f"network unavailable: {result['error']}")
        assert calls.count("example.com") == 1


class TestHtmlToText:
    def _text(self, html):
        from plugins.fetch_url import html_to_text
        return html_to_text(html)

    def test_keeps_text_and_title_drops_scripts_and_styles(self):
        title, text = self._text(
            "<html><head><title> My  Page </title><style>p{color:red}</style></head>"
            "<body><script>var x = 1;</script><p>Hello &amp; welcome</p><p>Second</p></body></html>"
        )
        assert title == "My Page"
        assert text == "Hello & welcome\nSecond"

    def test_drops_page_furniture_by_tag_role_class_and_hidden(self):
        _, text = self._text(
            "<body><nav>Menu</nav><div role='navigation'>Links</div><div class='cookie-banner'>Cookies</div>"
            "<div hidden>Secret</div><span aria-hidden='true'>Icon</span><footer>Footer</footer>"
            "<p>Real content</p></body>"
        )
        assert text == "Real content"

    def test_prefers_main_content_when_marked(self):
        _, text = self._text("<body><div>Banner outside</div><main><h1>Story</h1><p>Body</p></main><div>After</div></body>")
        assert text == "Story\nBody"

    def test_uses_whole_page_when_no_main(self):
        _, text = self._text("<body><div>One</div><div>Two</div></body>")
        assert text == "One\nTwo"

    def test_nested_same_tag_inside_skipped_region_does_not_end_it_early(self):
        _, text = self._text("<div class='sidebar'><div>inner</div>still sidebar</div><p>content</p>")
        assert text == "content"

    def test_list_items_get_bullets(self):
        _, text = self._text("<ul><li>a</li><li>b</li></ul>")
        assert text == "- a\n- b"

    def test_only_the_first_title_counts(self):
        title, _ = self._text("<head><title>Real</title></head><body><svg><title>icon</title></svg><p>x</p></body>")
        assert title == "Real"

    def test_word_matching_avoids_false_positives(self):
        # "shared-results" contains "share" but not as a whole word.
        _, text = self._text("<div class='shared-results'>keep me</div>")
        assert text == "keep me"


class TestFetchUrlContentTypes:
    def _fetch(self, monkeypatch, content_type, body):
        import httpx
        import plugins.fetch_url as fetch_module

        monkeypatch.setattr(fetch_module, "_resolve_pinned_ip", lambda host: ("93.184.215.14", None))

        def fake_send(self, request, **kwargs):
            headers = {"content-type": content_type} if content_type else {}
            return httpx.Response(200, headers=headers, content=body.encode(), request=request)

        monkeypatch.setattr(httpx.Client, "send", fake_send)
        return fetch_module.fetch_url("https://example.org/page")

    def test_html_is_converted_to_text(self, monkeypatch):
        result = self._fetch(monkeypatch, "text/html; charset=utf-8", "<title>T</title><main><p>Hi</p></main>")
        assert result == {"status": 200, "title": "T", "content": "Hi", "truncated": False}

    def test_json_and_plain_text_pass_through(self, monkeypatch):
        assert self._fetch(monkeypatch, "application/json", '{"a": 1}')["content"] == '{"a": 1}'
        assert self._fetch(monkeypatch, "text/plain", "plain <b>text</b>")["content"] == "plain <b>text</b>"

    def test_binary_content_is_refused(self, monkeypatch):
        result = self._fetch(monkeypatch, "application/pdf", "%PDF-1.7")
        assert "not a text page" in result["error"]

    def test_long_text_is_truncated(self, monkeypatch):
        import plugins.fetch_url as fetch_module
        monkeypatch.setattr(fetch_module, "MAX_CHARS", 10)
        result = self._fetch(monkeypatch, "text/plain", "x" * 50)
        assert result["content"] == "x" * 10
        assert result["truncated"] is True


class TestWebSearch:
    def test_missing_api_key_reports_error(self, monkeypatch):
        monkeypatch.delenv("BRAVE_API_KEY", raising=False)
        from plugins.web_search import web_search
        result = web_search("test query")
        assert "error" in result
        assert "BRAVE_API_KEY" in result["error"]


class TestShellExec:
    def _reload(self, monkeypatch, enabled, files_dir):
        if enabled:
            monkeypatch.setenv("AGENTIC_ENABLE_SHELL", "1")
        else:
            monkeypatch.delenv("AGENTIC_ENABLE_SHELL", raising=False)
        monkeypatch.setenv("AGENTIC_FILES_DIR", str(files_dir))
        import plugins.shell_exec as shell_exec
        return importlib.reload(shell_exec)

    def test_disabled_by_default(self, monkeypatch, tmp_path):
        module = self._reload(monkeypatch, enabled=False, files_dir=tmp_path)
        registered = {}

        class FakeRegistry:
            def register(self, name, func, spec):
                registered[name] = func

        module.register(FakeRegistry())
        assert registered == {}

    def test_enabled_via_env_var(self, monkeypatch, tmp_path):
        module = self._reload(monkeypatch, enabled=True, files_dir=tmp_path)
        registered = {}

        class FakeRegistry:
            def register(self, name, func, spec):
                registered[name] = func

        module.register(FakeRegistry())
        assert "run_command" in registered
        result = registered["run_command"]("echo hi")
        assert result["exit_code"] == 0
        assert "hi" in result["stdout"]

    def test_empty_command_rejected(self, monkeypatch, tmp_path):
        module = self._reload(monkeypatch, enabled=True, files_dir=tmp_path)
        result = module.run_command("   ")
        assert "error" in result

    def test_unknown_command_reports_error(self, monkeypatch, tmp_path):
        module = self._reload(monkeypatch, enabled=True, files_dir=tmp_path)
        result = module.run_command("this-command-should-not-exist-xyz")
        assert "error" in result


class TestWebSearchResults:
    def test_snippets_are_flattened_to_plain_text(self, monkeypatch):
        import httpx
        import plugins.web_search as web_search_module

        monkeypatch.setenv("BRAVE_API_KEY", "test-key")

        def fake_get(url, params=None, headers=None, timeout=None):
            assert headers["X-Subscription-Token"] == "test-key"
            return httpx.Response(200, json={"web": {"results": [{
                "title": "Signal &amp; privacy",
                "url": "https://example.org",
                "description": "The <strong>Signal</strong> app is <strong>encrypted</strong>.",
            }]}}, request=httpx.Request("GET", url))

        monkeypatch.setattr(web_search_module.httpx, "get", fake_get)
        result = web_search_module.web_search("signal")
        assert result == {"results": [{
            "title": "Signal & privacy", "url": "https://example.org", "snippet": "The Signal app is encrypted.",
        }]}
